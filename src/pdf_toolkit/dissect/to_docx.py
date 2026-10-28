"""Export reconstructed pages to DOCX, optionally filled.

Ported from ss-documentDissection's ``to_docx.py``. Each page becomes
its rendered image followed by a two-column table of field names and
``{{ name }}`` placeholders, which :func:`fill_docx_template` renders
with docxtpl. Needs python-docx and docxtpl (``pdf-toolkit[dissect]``
for the former via the office extra, docxtpl via the dissect extra).
"""

from __future__ import annotations

import io
import json
from pathlib import Path

from .models import FieldBox

__all__ = ["DocxError", "export_document_to_docx", "fill_docx_template", "page_to_docx"]


class DocxError(Exception):
    pass


def _require_docx():
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:
        raise DocxError(
            'DOCX export needs python-docx. Install it with: pip install "pdf-toolkit[dissect]"'
        ) from exc
    return Document, Inches


def _image_bytes_and_size(image, dpi: int):
    """Return ``(png_bytes, width_inches)`` for a NumPy image or a path."""
    from .detect import require_cv2

    cv2, np = require_cv2()
    if isinstance(image, (str, Path)):
        img_bytes = Path(image).read_bytes()
    else:
        success, buf = cv2.imencode(".png", image)
        if not success:
            raise DocxError("Failed to encode image")
        img_bytes = buf.tobytes()
    arr = cv2.imdecode(np.frombuffer(img_bytes, np.uint8), cv2.IMREAD_COLOR)
    h_px, w_px = arr.shape[:2]
    return img_bytes, w_px / dpi


def _add_field_table(doc, fields: list[FieldBox]) -> None:
    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text = "field name"
    hdr[1].text = "placeholder"
    for i, f in enumerate(fields):
        name = f.name or f"field_{i}"
        row = table.add_row().cells
        row[0].text = name
        row[1].text = f"{{{{ {name} }}}}"


def page_to_docx(image, fields: list[FieldBox], out_path: str | Path, dpi: int = 150) -> None:
    """Write one page (image + field table) to a .docx file."""
    Document, Inches = _require_docx()
    doc = Document()
    img_bytes, width_in = _image_bytes_and_size(image, dpi)
    doc.add_picture(io.BytesIO(img_bytes), width=Inches(width_in))
    doc.add_paragraph("")
    _add_field_table(doc, fields)
    path = Path(out_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def fill_docx_template(template_path: str | Path, data: dict, out_path: str | Path) -> None:
    """Render a docxtpl template (``{{ name }}`` placeholders) with *data*."""
    try:
        from docxtpl import DocxTemplate
    except ImportError as exc:
        raise DocxError(
            'DOCX filling needs docxtpl. Install it with: pip install "pdf-toolkit[dissect]"'
        ) from exc
    tpl = DocxTemplate(str(template_path))
    tpl.render(data)
    tpl.save(str(out_path))


def export_document_to_docx(
    doc_dir: str | Path, out_path: str | Path, fill_data: dict | None = None, dpi: int = 150
) -> None:
    """Export a reconstructed document directory (``images/`` + ``html/``) to one .docx.

    Pages are the ``page_*.png`` files in ``images/``, in filename
    order, each followed by its fieldmap table from ``html/``. If
    *fill_data* is given, the result is rendered through docxtpl in
    place, as in the original.
    """
    Document, Inches = _require_docx()
    base = Path(doc_dir)
    images_dir = base / "images"
    html_dir = base / "html"
    if not images_dir.exists():
        raise DocxError("images directory not found in doc_dir")
    doc = Document()
    for img in sorted(images_dir.glob("page_*.png")):
        img_bytes, width_in = _image_bytes_and_size(img, dpi)
        doc.add_picture(io.BytesIO(img_bytes), width=Inches(width_in))
        doc.add_paragraph("")
        page_num = None
        try:
            page_num = int("".join(ch for ch in img.stem if ch.isdigit()))
        except Exception:
            pass
        fields: list[FieldBox] = []
        if page_num is not None:
            fmap = html_dir / f"page_{page_num}_fieldmap.json"
            if fmap.exists():
                try:
                    raw = json.loads(fmap.read_text(encoding="utf-8"))
                    fields = [FieldBox.from_dict(item) for item in raw]
                except Exception:
                    fields = []
        _add_field_table(doc, fields)
        doc.add_page_break()
    path = Path(out_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    if fill_data:
        fill_docx_template(str(path), fill_data, str(path))
