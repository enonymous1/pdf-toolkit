"""Extraction and CLI tests for the citations checker.

Fixtures are built in-test with fictional text (the original's test
documents were a real RFP and a real proposal checklist and are not
shipped). Skips when python-docx is not installed, like the
formatting-checker tests.
"""

import json
from pathlib import Path

import fitz
import pytest

from pdf_toolkit.citations import extract_chunks_from_docx, extract_pages_from_pdf
from pdf_toolkit.cli import main as cli_main

docx = pytest.importorskip("docx")

VERBATIM = "The harbor lighthouse was built in eighteen forty two."


def _make_docx(path: Path, paragraphs=(), table_cells=()) -> Path:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_cells:
        table = doc.add_table(rows=1, cols=len(table_cells))
        for index, text in enumerate(table_cells):
            table.cell(0, index).text = text
    doc.save(str(path))
    return path


def _make_pdf(path: Path, page_texts) -> Path:
    doc = fitz.open()
    for text in page_texts:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=12)
    doc.save(str(path))
    doc.close()
    return path


def test_extract_chunks_includes_table_cells_in_order(tmp_path):
    path = _make_docx(
        tmp_path / "manuscript.docx",
        paragraphs=["Opening paragraph about the harbor."],
        table_cells=[VERBATIM, "Cell with unrelated fictional notes."],
    )
    chunks = extract_chunks_from_docx(path)
    texts = [chunk.text for chunk in chunks]
    assert texts[0] == "Opening paragraph about the harbor."
    assert VERBATIM in texts
    assert any(chunk.source.startswith("table 0") for chunk in chunks if chunk.text == VERBATIM)


def test_extract_chunks_splits_long_paragraphs(tmp_path):
    long_text = "First fictional sentence about the harbor. " + (
        "Second fictional sentence about the lighthouse keepers. " * 8
    )
    path = _make_docx(tmp_path / "long.docx", paragraphs=[long_text])
    chunks = extract_chunks_from_docx(path)
    assert len(chunks) > 1
    assert all(chunk.text for chunk in chunks)
    assert chunks[0].id == 0


def test_extract_pages_from_pdf(tmp_path):
    path = _make_pdf(tmp_path / "source.pdf", [VERBATIM, "Second fictional page."])
    pages = extract_pages_from_pdf(path)
    assert [page.page_number for page in pages] == [1, 2]
    assert VERBATIM in pages[0].text


def test_extract_missing_files_raise_file_not_found(tmp_path):
    with pytest.raises(FileNotFoundError):
        extract_chunks_from_docx(tmp_path / "missing.docx")
    with pytest.raises(FileNotFoundError):
        extract_pages_from_pdf(tmp_path / "missing.pdf")


def test_cli_end_to_end_json_report(tmp_path, capsys):
    docx_path = _make_docx(tmp_path / "manuscript.docx", paragraphs=[VERBATIM])
    pdf_path = _make_pdf(tmp_path / "source.pdf", [VERBATIM + " Extra source context."])
    report = tmp_path / "report.json"
    code = cli_main(["check-citations", str(docx_path), str(pdf_path), "--json-out", str(report)])
    assert code == 1
    data = json.loads(report.read_text(encoding="utf-8"))
    assert len(data) == 1
    assert data[0]["score"] == 100
    assert data[0]["page_number"] == 1
    assert VERBATIM in data[0]["pdf_snippet"]
    assert "Found 1 matches" in capsys.readouterr().out


def test_cli_no_matches_exits_zero(tmp_path, capsys):
    docx_path = _make_docx(
        tmp_path / "manuscript.docx",
        paragraphs=["Harbor tides follow a predictable lunar cycle each month."],
    )
    pdf_path = _make_pdf(
        tmp_path / "source.pdf", ["Desert rainfall is scarce and highly seasonal."]
    )
    code = cli_main(["check-citations", str(docx_path), str(pdf_path)])
    assert code == 0
    assert "No matches found." in capsys.readouterr().out


def test_cli_footnote_fragment_fixture_finds_nothing(tmp_path):
    # The original's own sample shape: checklist-style chunks against
    # a source full of footnote fragments. Must find nothing.
    docx_path = _make_docx(
        tmp_path / "checklist.docx",
        paragraphs=[
            "Must follow the specific format provided in the guide.",
            "Describe the qualifications and experience relevant here.",
        ],
    )
    pdf_path = _make_pdf(tmp_path / "source.pdf", ["1. P. 2. l. Unrelated source text."])
    code = cli_main(["check-citations", str(docx_path), str(pdf_path), "--threshold", "50"])
    assert code == 0


def test_cli_missing_file_exits_two(tmp_path, capsys):
    code = cli_main(["check-citations", str(tmp_path / "nope.docx"), str(tmp_path / "nope.pdf")])
    assert code == 2
    assert "Error:" in capsys.readouterr().out


def test_cli_threshold_out_of_range_exits_two(tmp_path, capsys):
    docx_path = _make_docx(tmp_path / "manuscript.docx", paragraphs=[VERBATIM])
    pdf_path = _make_pdf(tmp_path / "source.pdf", [VERBATIM])
    code = cli_main(["check-citations", str(docx_path), str(pdf_path), "--threshold", "999"])
    assert code == 2
    assert "Error:" in capsys.readouterr().out


def test_cli_skip_cited(tmp_path, capsys):
    cited = "The harbor lighthouse was built in eighteen forty two (Rivera, 2020)."
    docx_path = _make_docx(tmp_path / "manuscript.docx", paragraphs=[cited])
    pdf_path = _make_pdf(tmp_path / "source.pdf", [VERBATIM])
    assert cli_main(["check-citations", str(docx_path), str(pdf_path)]) == 1
    assert cli_main(["check-citations", str(docx_path), str(pdf_path), "--skip-cited"]) == 0
    capsys.readouterr()
