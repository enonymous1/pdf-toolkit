"""Tests for the ss-documentDissection fold-in: ``pdf_toolkit.dissect``.

Image/OCR/DOCX/server tests skip when the ``[dissect]`` extra is not
installed, mirroring ``test_to_pdf.py``; the keyword, fieldmap, and
page-extraction tests need only the core dependencies and always run.
"""

import importlib.util
import io
import json
import shutil
from pathlib import Path

import fitz
import pytest

from pdf_toolkit.cli import main as cli_main
from pdf_toolkit.dissect import (
    DetectionError,
    FieldBox,
    FieldmapError,
    FillError,
    extract_pages_to_bytes,
    keyword_match,
    load_fieldmap,
    page_texts,
    save_fieldmap,
    split_pdf_to_pages,
)
from pdf_toolkit.dissect.models import PageData

HAS_CV2 = importlib.util.find_spec("cv2") is not None
HAS_REPORTLAB = importlib.util.find_spec("reportlab") is not None
HAS_DOCX = importlib.util.find_spec("docx") is not None
HAS_FLASK = importlib.util.find_spec("flask") is not None

needs_cv2 = pytest.mark.skipif(not HAS_CV2, reason="OpenCV not installed (pdf-toolkit[dissect])")
needs_reportlab = pytest.mark.skipif(
    not HAS_REPORTLAB, reason="ReportLab not installed (pdf-toolkit[dissect])"
)
needs_docx = pytest.mark.skipif(
    not HAS_DOCX, reason="python-docx not installed (pdf-toolkit[dissect])"
)


def _make_pdf(path: Path, pages: list[str]) -> Path:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    doc.save(str(path))
    doc.close()
    return path


def _synthetic_form_image():
    import cv2
    import numpy as np

    img = np.full((400, 600, 3), 255, dtype=np.uint8)
    cv2.rectangle(img, (50, 50), (350, 80), (0, 0, 0), 2)
    cv2.rectangle(img, (50, 120), (550, 160), (0, 0, 0), 2)
    return img


def test_keyword_match_and_ratio():
    pages = [
        PageData(page_number=1, text="This is a form with name and date"),
        PageData(page_number=2, text="Other random page content"),
    ]
    res = keyword_match(pages, ["name", "date"], min_ratio=0.5)
    assert 1 in res.matched_pages
    assert 2 not in res.matched_pages
    assert res.scores[1] == 1.0
    res_strict = keyword_match(pages, ["name", "date", "signature"], min_ratio=1.0)
    assert res_strict.matched_pages == []


def test_keyword_match_empty_raises():
    with pytest.raises(DetectionError):
        keyword_match([PageData(page_number=1, text="x")], [])


def test_page_texts_and_extract(tmp_path: Path):
    pdf = _make_pdf(tmp_path / "doc.pdf", ["Page one alpha", "Page two beta", "Page three"])
    pages = page_texts(pdf)
    assert [p.page_number for p in pages] == [1, 2, 3]
    assert "alpha" in pages[0].text
    one = extract_pages_to_bytes(pdf, [2])
    doc = fitz.open(stream=one, filetype="pdf")
    assert doc.page_count == 1
    assert "beta" in doc.load_page(0).get_text("text")
    assert len(split_pdf_to_pages(pdf)) == 3


def test_find_template_pages_missing_pdf_raises(tmp_path: Path):
    from pdf_toolkit.dissect import find_template_pages

    with pytest.raises(DetectionError):
        find_template_pages(tmp_path / "missing.pdf", ["anything"])


def test_fieldmap_json_roundtrip(tmp_path: Path):
    fields = [
        FieldBox(name="name", x=10, y=20, w=100, h=30, score=0.9),
        FieldBox(x=1, y=2, w=3, h=4),
    ]
    path = save_fieldmap(fields, tmp_path / "page_1_fieldmap.json")
    # The file itself must be JSON: the original wrote str() and needed eval().
    assert json.loads(path.read_text(encoding="utf-8"))[0]["name"] == "name"
    loaded = load_fieldmap(path)
    assert loaded == fields


def test_fieldmap_repr_is_not_valid_json(tmp_path: Path):
    path = tmp_path / "bad.json"
    path.write_text("[{'name': 'x', 'x': 1, 'y': 2, 'w': 3, 'h': 4, 'score': None}]")
    with pytest.raises(FieldmapError):
        load_fieldmap(path)


def test_fieldmap_must_be_list(tmp_path: Path):
    path = tmp_path / "obj.json"
    path.write_text(json.dumps({"name": "x"}))
    with pytest.raises(FieldmapError):
        load_fieldmap(path)


@needs_cv2
def test_ssim_identical_images():
    import numpy as np

    from pdf_toolkit.dissect.detect import _ssim_gray

    a = np.full((200, 200), 128, dtype=np.uint8)
    assert 0.999 < _ssim_gray(a, a.copy()) <= 1.0


@needs_cv2
def test_image_template_match_identical_page(tmp_path: Path):
    import cv2

    from pdf_toolkit.dissect import image_template_match

    pdf = _make_pdf(tmp_path / "doc.pdf", ["Form Title"])
    doc = fitz.open(str(pdf))
    pix = doc.load_page(0).get_pixmap(matrix=fitz.Matrix(150 / 72, 150 / 72), alpha=False)
    pix.save(str(tmp_path / "template.png"))
    doc.close()
    res = image_template_match(pdf, tmp_path / "template.png", threshold=0.7, dpi=150)
    assert res.matched_pages == [1]
    assert cv2.imread is not None  # cv2 import exercised by the matcher


@needs_cv2
def test_detect_field_contours():
    from pdf_toolkit.dissect import detect_field_contours

    boxes = detect_field_contours(_synthetic_form_image(), min_area=100)
    assert len(boxes) >= 2
    assert boxes[0].y <= boxes[-1].y  # sorted top-to-bottom


@needs_cv2
def test_create_html_overlay_escapes_values():
    from pdf_toolkit.dissect import create_html_overlay

    html_text = create_html_overlay(
        _synthetic_form_image(),
        [FieldBox(name='a"b', x=10, y=10, w=50, h=20)],
        data={'a"b': "<script>alert(1)</script>"},
    )
    assert "<script>alert(1)</script>" not in html_text
    assert "&lt;script&gt;" in html_text
    assert "position:absolute" in html_text


@needs_cv2
@needs_reportlab
def test_fillable_pdf_fill_roundtrip():
    from pypdf import PdfReader

    from pdf_toolkit.dissect import create_fillable_pdf_from_image, fill_pdf_fields

    img = _synthetic_form_image()
    boxes = [FieldBox(name="field_0", x=50, y=50, w=300, h=30)]
    pdf_bytes = create_fillable_pdf_from_image(img, boxes, dpi=150)
    assert pdf_bytes.startswith(b"%PDF")
    filled = fill_pdf_fields(pdf_bytes, {"field_0": "Alice"}, fields=boxes, dpi=150)
    fields = PdfReader(io.BytesIO(filled)).get_fields() or {}
    assert fields["field_0"].get("/V") == "Alice"
    doc = fitz.open(stream=filled, filetype="pdf")
    assert "Alice" in doc.load_page(0).get_text("text")


@needs_reportlab
def test_fill_multipage_acroform():
    from pypdf import PdfReader
    from reportlab.pdfgen import canvas

    from pdf_toolkit.dissect import fill_pdf_fields

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.acroForm.textfield(name="field_a", x=72, y=700, width=200, height=20)
    c.showPage()
    c.acroForm.textfield(name="field_b", x=72, y=700, width=200, height=20)
    c.showPage()
    c.save()
    filled = fill_pdf_fields(buf.getvalue(), {"field_a": "Alice", "field_b": "Bob"})
    fields = PdfReader(io.BytesIO(filled)).get_fields() or {}
    assert fields["field_a"].get("/V") == "Alice"
    assert fields["field_b"].get("/V") == "Bob"


def test_fill_overlay_fallback_and_path_rejection(tmp_path: Path):
    from pdf_toolkit.dissect import fill_pdf_fields

    pdf = _make_pdf(tmp_path / "plain.pdf", ["plain page"])
    plain = pdf.read_bytes()
    boxes = {1: [FieldBox(name="field_0", x=50, y=50, w=300, h=40)]}
    filled = fill_pdf_fields(plain, {"field_0": "Hello"}, fields=boxes, dpi=150)
    doc = fitz.open(stream=filled, filetype="pdf")
    assert "Hello" in doc.load_page(0).get_text("text")
    with pytest.raises(FillError):
        fill_pdf_fields(str(pdf), {"field_0": "Hello"})
    with pytest.raises(FillError):
        fill_pdf_fields(plain, {"field_0": "Hello"})


def test_fill_overlay_shrinks_font_for_short_boxes(tmp_path: Path):
    """A 23px-tall box at 150dpi is 11pt: a fixed 10pt font inserts nothing."""
    from pdf_toolkit.dissect import fill_pdf_fields

    pdf = _make_pdf(tmp_path / "plain.pdf", ["plain page"])
    boxes = {1: [FieldBox(name="field_0", x=210, y=133, w=215, h=23)]}
    filled = fill_pdf_fields(pdf.read_bytes(), {"field_0": "Alice"}, fields=boxes, dpi=150)
    doc = fitz.open(stream=filled, filetype="pdf")
    assert "Alice" in doc.load_page(0).get_text("text")


@needs_cv2
@needs_reportlab
def test_reconstruct_end_to_end(tmp_path: Path):
    from pdf_toolkit.dissect import reconstruct_pdf_pages

    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 72), "Form Title Name Signature")
    page.draw_rect(fitz.Rect(72, 120, 400, 150), color=(0, 0, 0), width=2)
    page.draw_rect(fitz.Rect(72, 180, 400, 210), color=(0, 0, 0), width=2)
    pdf = tmp_path / "form.pdf"
    doc.save(str(pdf))
    doc.close()
    out = tmp_path / "recon"
    results = reconstruct_pdf_pages(pdf, [1], out, dpi=150)
    info = results[1]
    assert info["fields_detected"] >= 1
    for key in ("image", "html", "fieldmap", "fillable_pdf", "extracted_pdf"):
        assert Path(info[key]).exists()
    # Fieldmap is JSON, loadable without eval.
    fieldmap = json.loads(Path(info["fieldmap"]).read_text(encoding="utf-8"))
    assert isinstance(fieldmap, list) and len(fieldmap) == info["fields_detected"]
    assert "<input" in Path(info["html"]).read_text(encoding="utf-8")
    assert Path(info["fillable_pdf"]).read_bytes().startswith(b"%PDF")
    extracted = fitz.open(stream=Path(info["extracted_pdf"]).read_bytes(), filetype="pdf")
    assert extracted.page_count == 1


@needs_cv2
@needs_docx
def test_page_to_docx_and_fill(tmp_path: Path):
    from docx import Document

    from pdf_toolkit.dissect import fill_docx_template, page_to_docx

    out = tmp_path / "page_1.docx"
    page_to_docx(_synthetic_form_image(), [FieldBox(name="name", x=10, y=10, w=50, h=20)], str(out))
    text = "\n".join(c.text for t in Document(str(out)).tables for r in t.rows for c in r.cells)
    assert "{{ name }}" in text
    filled = tmp_path / "filled.docx"
    fill_docx_template(str(out), {"name": "Alice"}, str(filled))
    filled_text = "\n".join(
        c.text for t in Document(str(filled)).tables for r in t.rows for c in r.cells
    )
    assert "Alice" in filled_text


@pytest.mark.skipif(
    shutil.which("tesseract") is not None, reason="Tesseract present; labelling may succeed"
)
@needs_cv2
def test_ocr_labelling_degrades_gracefully_without_tesseract():
    from pdf_toolkit.dissect import assign_labels_to_fields

    fields = [FieldBox(x=50, y=50, w=300, h=30)]
    result = assign_labels_to_fields(_synthetic_form_image(), fields)
    assert result[0].name is None


def test_cli_dissect_detect_keyword(tmp_path: Path, capsys):
    pdf = _make_pdf(tmp_path / "doc.pdf", ["Form Title with name and date", "unrelated"])
    assert cli_main(["dissect", "detect", str(pdf), "--keyword", "name", "--keyword", "date"]) == 0
    out = capsys.readouterr().out
    assert '"matched_pages": [\n    1' in out or '"matched_pages": [1]' in out.replace("\n", "")


def test_cli_dissect_missing_pdf_returns_2(tmp_path: Path):
    assert cli_main(["dissect", "detect", str(tmp_path / "nope.pdf"), "--keyword", "x"]) == 2


@needs_cv2
@needs_reportlab
def test_cli_dissect_reconstruct_fill_export_docx(tmp_path: Path):
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 72), "Form Title Name Signature")
    page.draw_rect(fitz.Rect(72, 120, 400, 150), color=(0, 0, 0), width=2)
    pdf = tmp_path / "form.pdf"
    doc.save(str(pdf))
    doc.close()
    recon = tmp_path / "recon"
    assert (
        cli_main(
            [
                "dissect",
                "reconstruct",
                str(pdf),
                "--keyword",
                "Form",
                "--keyword",
                "Title",
                "--out-dir",
                str(recon),
            ]
        )
        == 0
    )
    assert (recon / "html" / "page_1_fieldmap.json").exists()
    data = tmp_path / "data.json"
    data.write_text(json.dumps({"1": {"field_0": "Alice"}}))
    filled_dir = tmp_path / "filled"
    assert (
        cli_main(
            [
                "dissect",
                "fill",
                str(pdf),
                "--fieldmap",
                str(recon),
                "--data",
                str(data),
                "--out-dir",
                str(filled_dir),
            ]
        )
        == 0
    )
    assert (filled_dir / "page_1_filled.pdf").exists()
    docx_dir = tmp_path / "docx"
    assert (
        cli_main(
            [
                "dissect",
                "export-docx",
                "--fieldmap",
                str(recon),
                "--out-dir",
                str(docx_dir),
                "--fill",
                "--data",
                str(data),
            ]
        )
        == 0
    )
    assert (docx_dir / "page_1.docx").exists()
    assert (docx_dir / "page_1_filled.docx").exists()


@pytest.mark.skipif(not HAS_FLASK, reason="Flask not installed (pdf-toolkit[dissect])")
@needs_cv2
@needs_reportlab
def test_server_confined_to_root(tmp_path: Path):
    import json as _json

    from pdf_toolkit.dissect import create_app, reconstruct_pdf_pages

    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 72), "Form Title")
    page.draw_rect(fitz.Rect(72, 120, 400, 150), color=(0, 0, 0), width=2)
    pdf = tmp_path / "form.pdf"
    doc.save(str(pdf))
    doc.close()
    root = tmp_path / "recon"
    reconstruct_pdf_pages(pdf, [1], root, dpi=150)
    secret = tmp_path / "secret.txt"
    secret.write_text("top secret")
    client = create_app(root).test_client()

    assert client.get("/index").status_code == 200
    assert client.get("/files/../secret.txt").status_code == 404
    assert client.get("/files/html/page_1.html").status_code == 200
    assert client.get("/thumbnail?path=images/page_1.png&size=0").status_code == 400
    assert client.get("/thumbnail?path=images/page_1.png&size=80").status_code == 200

    # Absolute and escaping fieldmap writes are rejected.
    rv = client.post(
        "/save_fieldmap",
        json={"fieldmap": [], "out_path": str(tmp_path / "evil.json")},
    )
    assert rv.status_code == 400
    assert not (tmp_path / "evil.json").exists()
    rv = client.post("/save_fieldmap", json={"fieldmap": [], "out_path": "../evil.json"})
    assert rv.status_code == 400
    # A write inside the root succeeds and is real JSON.
    rv = client.post(
        "/save_fieldmap",
        json={
            "fieldmap": [{"name": "field_0", "x": 1, "y": 2, "w": 3, "h": 4}],
            "out_path": "html/page_1_fieldmap.json",
        },
    )
    assert rv.status_code == 200
    assert _json.loads((root / "html" / "page_1_fieldmap.json").read_text())[0]["name"] == "field_0"

    # Fill-document now fills and writes a file (the original always errored).
    rv = client.post("/fill-document", json={"group_dir": ".", "data": {"field_0": "Alice"}})
    assert rv.status_code == 200
    filled = rv.get_json()["filled_files"]
    assert filled and all(isinstance(f, str) for f in filled)
    assert (root / filled[0]).exists()
